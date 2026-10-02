from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches
import os
import time

def generate():
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        # Using 16:9 ratio, 1920x1080
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        html_path = os.path.abspath('pitch_deck.html')
        page.goto(f"file:///{html_path}")
        
        # Hide UI controls
        page.evaluate("if(document.getElementById('controls')) document.getElementById('controls').style.display = 'none';")
        page.evaluate("if(document.getElementById('progress-bar-wrap')) document.getElementById('progress-bar-wrap').style.display = 'none';")

        prs = Presentation()
        # Set 16:9 Aspect Ratio for PPTX
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        for i in range(1, 13):
            print(f"Capturing slide {i}...")
            # Deactivate all
            page.evaluate("document.querySelectorAll('.slide').forEach(s => { s.classList.remove('active'); s.style.display = 'none'; s.style.opacity = '0'; });")
            # Activate current
            script = f"""
                let s = document.getElementById('slide-{i}');
                if (s) {{
                    s.classList.add('active');
                    s.style.display = 'flex';
                    s.style.opacity = '1';
                }}
            """
            page.evaluate(script)
            
            # Wait for any CSS animations to finish
            page.wait_for_timeout(1500)
            
            img_path = f"slide_{i}.png"
            page.screenshot(path=img_path)
            
            # Add to PPTX
            blank_slide = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_slide)
            slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
            
        prs.save("InsightO_Pitch_Deck.pptx")
        browser.close()
        
        # Cleanup images
        for i in range(1, 13):
            try:
                os.remove(f"slide_{i}.png")
            except:
                pass
        
        print("PPTX generated successfully as InsightO_Pitch_Deck.pptx")

if __name__ == "__main__":
    generate()
